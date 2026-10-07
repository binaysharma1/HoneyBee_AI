"""Upload and extract basic facts from common document types."""

from __future__ import annotations

import csv
import io
import logging
import re
from typing import Any

logger = logging.getLogger("honeybee.documents")

_TABLE_COLS = 8


def extract_fact_text(filename: str, content: bytes) -> str:
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    try:
        if suffix in ("txt", "md"):
            return content.decode("utf-8", errors="replace")
        if suffix == "pdf":
            return _extract_pdf(content)
        if suffix == "docx":
            return _extract_docx(content)
        if suffix == "csv":
            return _summarize_csv(content)
        if suffix in ("xlsx", "xls"):
            return _summarize_excel(content)
        return content.decode("utf-8", errors="replace")
    except Exception as exc:
        logger.warning("extract failed for %s: %s", filename, exc)
        raise RuntimeError(f"Failed to parse {suffix} document: {exc}") from exc


def summarise_csv(text: str) -> str:
    lines = [ln for ln in text.splitlines() if ln.strip()]
    if not lines:
        return "Empty CSV."
    reader = csv.reader(lines[:40])
    rows = list(reader)
    header = rows[0] if rows else []
    stats = [f"Columns: {', '.join(header[:_TABLE_COLS])}{'...' if len(header) > _TABLE_COLS else ''}"]
    for row in rows[1:4]:
        cells = [f"{col}: {row[i] if i < len(row) else ''}" for i, col in enumerate(header[:_TABLE_COLS])]
        stats.append("Sample row: " + ", ".join(cells))
    stats.append(f"Rows analysed: {min(len(rows)-1, 1000)}")
    return "\n".join(stats)


def _summarize_csv(content: bytes) -> str:
    text = content.decode("utf-8", errors="replace")
    return summarise_csv(text)


def _summarize_excel(content: bytes) -> str:
    try:
        import pandas as pd
    except Exception as exc:
        return f"Excel reader unavailable: {exc}"
    df = pd.read_excel(io.BytesIO(content), sheet_name=None)
    parts = []
    for sheet_name, df_sheet in df.items():
        cols = ", ".join(list(df_sheet.columns[:_TABLE_COLS]))
        parts.append(f"Sheet '{sheet_name}': columns {cols}, rows: {len(df_sheet)}")
        preview = df_sheet.head(2).to_csv(index=False, header=True).strip().splitlines()
        if preview:
            parts.append("Sample:\n" + "\n".join(preview))
    return "\n\n".join(parts)


def _extract_pdf(content: bytes) -> str:
    try:
        from pypdf import PdfReader
    except Exception as exc:
        raise RuntimeError("pypdf not installed") from exc
    reader = PdfReader(io.BytesIO(content))
    pages = []
    for i, page in enumerate(reader.pages[:40], start=1):
        txt = page.extract_text() or ""
        pages.append(f"Page {i}:\n{txt.strip()[:800]}")
    return "\n\n".join(pages)


def _extract_docx(content: bytes) -> str:
    try:
        import docx
    except Exception as exc:
        raise RuntimeError("python-docx not installed") from exc
    doc = docx.Document(io.BytesIO(content))
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows[:5]:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(lines[:200])


def chunk_text(text: str, size: int = 900, overlap: int = 120) -> list[str]:
    words = re.split(r"\s+", text)
    if not words:
        return []
    chunks = []
    current: list[str] = []
    current_len = 0
    for word in words:
        current.append(word)
        current_len += len(word) + 1
        if current_len > size:
            chunks.append(" ".join(current))
            # start new chunk keeping overlap tail
            tail = " ".join(current)[-overlap:]
            current = [tail] if tail else []
            current_len = len(tail)
    if current:
        chunks.append(" ".join(current))
    return chunks


def extract_facts_for_file(filename: str, content: bytes) -> tuple[str, list[str]]:
    """Returns a summary fact and list of chunk facts for the file."""
    text = extract_fact_text(filename, content)
    chunks = chunk_text(text, size=900, overlap=120)
    prefix = f"Uploaded document: {filename}. Extract type: basic content facts."
    summary = text[:800].replace("\n", " ")
    return prefix + " " + summary, chunks
